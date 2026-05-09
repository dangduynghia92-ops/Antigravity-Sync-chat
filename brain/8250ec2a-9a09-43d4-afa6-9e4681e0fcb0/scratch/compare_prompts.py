import json, os

base_new = r'C:\Users\Admin\OneDrive\Documents\Biography_Saladin_(Ayyubid)_20260503_1700\v1_Cuộc_Đời_Bạn\Test 3\video_prompt'
base_old = r'C:\Users\Admin\OneDrive\Documents\Biography_Saladin_(Ayyubid)_20260503_1700\v1_Cuộc_Đời_Bạn\Test 3\video_prompt_v1'

def load_chars(base):
    f = os.path.join(base, 'ch_08_Level_8__Chilled_step2_characters.json')
    if os.path.exists(f):
        return json.load(open(f, encoding='utf-8')).get('characters', [])
    return []

chars_new = load_chars(base_new)
chars_old = load_chars(base_old)

print("=" * 80)
print("SO SÁNH PROMPT NHÂN VẬT CHÍNH (Step 2b)")
print("=" * 80)

# Build lookup by label
old_map = {c.get('label', ''): c for c in chars_old}
new_map = {c.get('label', ''): c for c in chars_new}

all_labels = sorted(set(list(old_map.keys()) + list(new_map.keys())))

for label in all_labels:
    old = old_map.get(label)
    new = new_map.get(label)
    
    print(f"\n{'─' * 80}")
    print(f"🏷  {label}")
    print(f"{'─' * 80}")
    
    if old and new:
        # Compare each field
        for field in ['original_name', 'visual_description', 'sheet_prompt', 'default_stance']:
            old_val = old.get(field, '')
            new_val = new.get(field, '')
            if old_val != new_val:
                print(f"\n  📌 {field}:")
                print(f"     OLD: {str(old_val)[:300]}")
                print(f"     NEW: {str(new_val)[:300]}")
            else:
                print(f"  ✅ {field}: SAME")
    elif old and not new:
        print(f"  ❌ MISSING in NEW")
    elif new and not old:
        print(f"  🆕 NEW only")

# Also compare crowd data
print(f"\n\n{'=' * 80}")
print("SO SÁNH CROWD (Step 2d)")
print("=" * 80)

def load_crowd(base):
    f = os.path.join(base, 'ch_08_Level_8__Chilled_step2d_crowd.json')
    if os.path.exists(f):
        return json.load(open(f, encoding='utf-8')).get('characters', [])
    return []

crowd_old = load_crowd(base_old)
crowd_new = load_crowd(base_new)

print(f"  OLD: {len(crowd_old)} crowd sheets")
print(f"  NEW: {len(crowd_new)} crowd sheets")

if crowd_old:
    print(f"\n  OLD crowd labels: {[c.get('label') for c in crowd_old]}")
if crowd_new:
    print(f"  NEW crowd labels: {[c.get('label') for c in crowd_new]}")

# Compare sheet_prompt format
if crowd_old:
    sp = crowd_old[0].get('sheet_prompt', '')
    print(f"\n  OLD crowd[0] sheet_prompt: {sp[:200]}")
if crowd_new:
    sp = crowd_new[0].get('sheet_prompt', '')
    print(f"  NEW crowd[0] sheet_prompt: {sp[:200]}")

# Also compare Excel prompt for named characters
print(f"\n\n{'=' * 80}")
print("SO SÁNH EXCEL PROMPT (Sheet 2 — cột F)")
print("=" * 80)

try:
    from openpyxl import load_workbook
    import glob
    
    for tag, base in [("OLD", base_old), ("NEW", base_new)]:
        xlsx = glob.glob(os.path.join(base, '*_video_prompts.xlsx'))
        if xlsx:
            wb = load_workbook(xlsx[0], read_only=True)
            ws = wb['Reference Images']
            print(f"\n  --- {tag} ({os.path.basename(xlsx[0])}) ---")
            print(f"  Total rows: {ws.max_row - 1}")
            for row in ws.iter_rows(min_row=2, max_row=min(6, ws.max_row), values_only=True):
                typ = row[1] if len(row) > 1 else ''
                label = row[2] if len(row) > 2 else ''
                prompt = str(row[5])[:150] if len(row) > 5 and row[5] else ''
                print(f"    [{typ}] {label}: {prompt}...")
            wb.close()
        else:
            print(f"\n  --- {tag}: No xlsx found ---")
except Exception as e:
    print(f"  Excel compare error: {e}")
