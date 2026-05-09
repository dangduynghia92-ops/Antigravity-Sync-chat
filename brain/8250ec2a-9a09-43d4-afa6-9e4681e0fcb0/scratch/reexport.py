"""Re-export Excel with crowd characters using current checkpoint data."""
import json, os, sys, re, glob

sys.path.insert(0, r'f:\1. Edit Videos\8.AntiCode\1.Prompt_Image\1.Prompt_Image')

base = os.path.join(r'C:\Users\Admin\OneDrive\Documents\Biography_Saladin_(Ayyubid)_20260503_1700',
    'v1_Cuộc_Đời_Bạn', 'Test 3', 'video_prompt')

# Load all data
prompts = json.load(open(os.path.join(base, 'ch_08_Level_8__Chilled_step4_prompts.json'), encoding='utf-8'))
chars_data = json.load(open(os.path.join(base, 'ch_08_Level_8__Chilled_step2_characters.json'), encoding='utf-8'))
crowd_data = json.load(open(os.path.join(base, 'ch_08_Level_8__Chilled_step2d_crowd.json'), encoding='utf-8'))

# Load style
style_path = os.path.join(r'f:\1. Edit Videos\8.AntiCode\1.Prompt_Image\1.Prompt_Image\video_styles', 'Chibi Storybook Historical.txt')
with open(style_path, encoding='utf-8') as f:
    sc = f.read()
ms = re.search(r'Mandatory\s+Style[:\s]*(.+?)(?:Negative|$)', sc, re.DOTALL | re.IGNORECASE)
mandatory_style = ms.group(1).strip() if ms else ""
np_m = re.search(r'Negative\s+Prompt[:\s]*(.+?)$', sc, re.DOTALL | re.IGNORECASE)
negative_prompt = np_m.group(1).strip() if np_m else ""

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

wb = Workbook()
header_fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
data_font = Font(name="Segoe UI", size=10)
wrap_align = Alignment(wrap_text=True, vertical="top")
center_align = Alignment(horizontal="center", vertical="top")
thin_border = Border(
    left=Side(style="thin", color="4B5563"), right=Side(style="thin", color="4B5563"),
    top=Side(style="thin", color="4B5563"), bottom=Side(style="thin", color="4B5563"))

def _write_headers(ws, headers):
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

# Sheet 1: Video Prompts
ws1 = wb.active
ws1.title = "Video Prompts"
_write_headers(ws1, ["#", "Chapter", "Scene ID", "Duration", "Characters", "Character Info", "Crowd", "Location", "Flat Prompt", "Negative Prompt"])

for row_idx, p in enumerate(prompts, 2):
    ws1.cell(row=row_idx, column=1, value=row_idx - 1).alignment = center_align
    ws1.cell(row=row_idx, column=2, value=p.get("chapter", "")).alignment = wrap_align
    ws1.cell(row=row_idx, column=3, value=p.get("global_scene_id", "")).alignment = center_align
    ws1.cell(row=row_idx, column=4, value=p.get("duration", "")).alignment = center_align
    ws1.cell(row=row_idx, column=5, value=p.get("characters", "")).alignment = wrap_align
    ws1.cell(row=row_idx, column=6, value=p.get("character_info", "")).alignment = wrap_align
    crowd_str = ", ".join(p.get("crowd_labels", [])) if p.get("crowd_labels") else ""
    ws1.cell(row=row_idx, column=7, value=crowd_str).alignment = wrap_align
    ws1.cell(row=row_idx, column=8, value=p.get("location", "")).alignment = wrap_align
    ws1.cell(row=row_idx, column=9, value=p.get("flat_prompt", "")).alignment = wrap_align
    ws1.cell(row=row_idx, column=10, value=p.get("negative_prompt", "")).alignment = wrap_align
    for col_idx in range(1, 11):
        ws1.cell(row=row_idx, column=col_idx).font = data_font
        ws1.cell(row=row_idx, column=col_idx).border = thin_border

ws1.column_dimensions["A"].width = 5
ws1.column_dimensions["B"].width = 25
ws1.column_dimensions["C"].width = 18
ws1.column_dimensions["D"].width = 9
ws1.column_dimensions["E"].width = 25
ws1.column_dimensions["F"].width = 40
ws1.column_dimensions["G"].width = 30
ws1.column_dimensions["H"].width = 20
ws1.column_dimensions["I"].width = 80
ws1.column_dimensions["J"].width = 40
ws1.freeze_panes = "A2"

# Sheet 2: Reference Images
ws2 = wb.create_sheet("Reference Images")
_write_headers(ws2, ["#", "Type", "Label", "Original Name", "Description", "Prompt (gen ảnh tham chiếu)", "Negative Prompt"])

row = 2
# Characters
for c in chars_data.get("characters", []):
    label = c.get("label", "")
    visual_desc = c.get("visual_description", "")
    sheet_prompt = c.get("sheet_prompt", "")
    if not sheet_prompt:
        sheet_prompt = f"Character reference sheet on clean white background. Bold rounded sans-serif text '{label}' at top. Three views: front, 3/4, side profile. Neutral standing pose. {visual_desc}."
    flat = f"{mandatory_style}. {sheet_prompt}" if mandatory_style else sheet_prompt

    ws2.cell(row=row, column=1, value=row - 1).alignment = center_align
    ws2.cell(row=row, column=2, value="Character").alignment = center_align
    ws2.cell(row=row, column=3, value=label).alignment = wrap_align
    ws2.cell(row=row, column=4, value=c.get("original_name", "")).alignment = wrap_align
    ws2.cell(row=row, column=5, value=visual_desc).alignment = wrap_align
    ws2.cell(row=row, column=6, value=flat).alignment = wrap_align
    ws2.cell(row=row, column=7, value=negative_prompt).alignment = wrap_align
    for col_idx in range(1, 8):
        ws2.cell(row=row, column=col_idx).font = data_font
        ws2.cell(row=row, column=col_idx).border = thin_border
    row += 1

# CROWD Characters
for c in crowd_data.get("characters", []):
    label = c.get("label", "")
    crowd_type = c.get("crowd_type", "")
    faction = c.get("faction", "")
    visual_desc = c.get("visual_description", "")
    sheet_prompt = c.get("sheet_prompt", "")
    if not sheet_prompt:
        sheet_prompt = f"Character reference sheet on clean white background. Full body front view of a {faction} {crowd_type}. {visual_desc}. Standing in {c.get('default_stance', 'neutral pose')}. Bold text label: {label}."
    flat = f"{mandatory_style}. {sheet_prompt}" if mandatory_style else sheet_prompt

    ws2.cell(row=row, column=1, value=row - 1).alignment = center_align
    ws2.cell(row=row, column=2, value="Crowd").alignment = center_align
    ws2.cell(row=row, column=3, value=label).alignment = wrap_align
    ws2.cell(row=row, column=4, value=f"{faction} — {crowd_type}").alignment = wrap_align
    ws2.cell(row=row, column=5, value=visual_desc).alignment = wrap_align
    ws2.cell(row=row, column=6, value=flat).alignment = wrap_align
    ws2.cell(row=row, column=7, value=negative_prompt).alignment = wrap_align
    for col_idx in range(1, 8):
        ws2.cell(row=row, column=col_idx).font = data_font
        ws2.cell(row=row, column=col_idx).border = thin_border
    row += 1

ws2.column_dimensions["A"].width = 5
ws2.column_dimensions["B"].width = 12
ws2.column_dimensions["C"].width = 30
ws2.column_dimensions["D"].width = 25
ws2.column_dimensions["E"].width = 40
ws2.column_dimensions["F"].width = 100
ws2.column_dimensions["G"].width = 40
ws2.freeze_panes = "A2"

# Save
xlsx_path = os.path.join(base, 'ch_08_Level_8__Chilled_video_prompts.xlsx')
wb.save(xlsx_path)
print(f"✅ Excel re-exported: {xlsx_path}")
print(f"   Sheet 1: {len(prompts)} prompts")
print(f"   Sheet 2: {len(chars_data.get('characters',[]))} characters + {len(crowd_data.get('characters',[]))} crowd")
